import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from src.extractors import (
    DocxExtractor,
    ExtractedDocument,
    ExtractorFactory,
    PDFExtractor,
    TextExtractor,
)


class TestExtractors(unittest.TestCase):

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_text_extractor_utf8(self):
        file_path = self.test_dir / "sample.txt"
        file_path.write_text("Hello World\nThis is a test file for RAG.", encoding="utf-8")

        extractor = TextExtractor()
        doc = extractor.extract(file_path)

        self.assertIsInstance(doc, ExtractedDocument)
        self.assertIn("Hello World", doc.text)
        self.assertEqual(doc.metadata["file_name"], "sample.txt")
        self.assertEqual(doc.metadata["line_count"], 2)

    def test_text_extractor_markdown(self):
        file_path = self.test_dir / "sample.md"
        file_path.write_text("# Title\n\nSome **markdown** notes.", encoding="utf-8")

        extractor = TextExtractor()
        doc = extractor.extract(file_path)

        self.assertIn("# Title", doc.text)
        self.assertIn("markdown", doc.text)

    def test_docx_extractor(self):
        file_path = self.test_dir / "sample.docx"
        doc = Document()
        doc.add_heading("RAG Architecture", level=1)
        doc.add_paragraph("This is the first paragraph.")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Header 1"
        table.rows[0].cells[1].text = "Header 2"
        table.rows[1].cells[0].text = "Value A"
        table.rows[1].cells[1].text = "Value B"
        doc.save(str(file_path))

        extractor = DocxExtractor()
        extracted = extractor.extract(file_path)

        self.assertIn("RAG Architecture", extracted.text)
        self.assertIn("first paragraph", extracted.text)
        self.assertIn("Header 1 | Header 2", extracted.text)
        self.assertEqual(extracted.metadata["table_count"], 1)

    def test_pdf_extractor(self):
        file_path = self.test_dir / "sample.pdf"
        writer = PdfWriter()
        # Create a blank page
        writer.add_blank_page(width=200, height=200)
        with open(file_path, "wb") as f:
            writer.write(f)

        extractor = PDFExtractor()
        extracted = extractor.extract(file_path)
        self.assertEqual(extracted.metadata["total_pages"], 1)

    def test_extractor_factory(self):
        self.assertTrue(ExtractorFactory.supports(Path("test.pdf")))
        self.assertTrue(ExtractorFactory.supports(Path("test.docx")))
        self.assertTrue(ExtractorFactory.supports(Path("test.txt")))
        self.assertTrue(ExtractorFactory.supports(Path("test.md")))
        self.assertFalse(ExtractorFactory.supports(Path("test.exe")))

        self.assertIsInstance(ExtractorFactory.get_extractor(Path("doc.pdf")), PDFExtractor)
        self.assertIsInstance(ExtractorFactory.get_extractor(Path("doc.docx")), DocxExtractor)
        self.assertIsInstance(ExtractorFactory.get_extractor(Path("doc.txt")), TextExtractor)

        with self.assertRaises(ValueError):
            ExtractorFactory.get_extractor(Path("doc.unsupported"))


if __name__ == "__main__":
    unittest.main()
