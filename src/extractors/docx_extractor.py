import logging
from pathlib import Path
from docx import Document

from src.extractors.base import BaseExtractor, ExtractedDocument

logger = logging.getLogger(__name__)


class DocxExtractor(BaseExtractor):
    """Extracts text, tables, and metadata from DOCX documents."""

    def extract(self, file_path: Path) -> ExtractedDocument:
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        try:
            doc = Document(str(file_path))
        except Exception as e:
            logger.error(f"Failed to read DOCX {file_path}: {e}")
            raise

        content_parts: list[str] = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content_parts.append(text)

        # Extract tables
        for table in doc.tables:
            table_rows: list[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_rows.append(" | ".join(row_cells))
            if table_rows:
                content_parts.append("\n".join(table_rows))

        full_text = "\n\n".join(content_parts)

        # Extract doc properties if present
        meta = {
            "file_name": file_path.name,
            "extension": file_path.suffix.lower(),
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "title": doc.core_properties.title or "",
            "author": doc.core_properties.author or "",
        }

        return ExtractedDocument(
            text=full_text,
            file_path=file_path,
            metadata=meta
        )
